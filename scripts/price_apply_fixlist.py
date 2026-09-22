#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
비라핑 오류 정정 — 「구간 밀림」 + 명시 지정 건만.

대상 (한경준님 확정 2026-09-22):
  ① 구간 밀림 — 현재값이 **같은 블록의 다른 직경 구간 값**과 정확히 일치하는 건.
     정가 변동이 아니라 기입 시 칸을 잘못 읽은 것 (wiki §5 「한 칸 위로 미는 실수」).
  ② EXTRA_CODES — 원인이 확정된 개별 오류 (자릿수 누락 등).

제외:
  · 1날·5날·10날 — 정가표에 그 날수 행이 없어 생기는 차이. 고치면 가격 정책 변경이 된다.
  · IGNORE_CODES — 한경준님이 직접 수정한 건.

⚠️ `_완료.xlsx` 에 **덮어쓴다**(라핑 정정 위에 얹는다). 원본 3파일은 건드리지 않는다.
   zip 을 메모리에서 완성한 뒤 한 번에 기록한다 — 중간 실패로 파일이 깨지지 않게.
"""
import argparse, io, os, re, sys, zipfile
import openpyxl
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from price_audit import PriceList, Blocks, parse_name, SHEET_ITEMS, ROW_START, \
                        COL_CODE, COL_NAME, COL_PRICE
from openpyxl.utils import get_column_letter
PC = get_column_letter(COL_PRICE)

EXTRA_CODES = {"3FL2501200"}          # 3날/평/25mm/초경/옆날/비코팅 — 2,660 = 정가의 1/10 (자릿수 누락)
IGNORE_CODES = {"2FL2201100", "K2FL2201100"}
SKIP_BLADES = {1, 5, 10}              # 정가표에 없는 날수 — 건드리지 않는다
PROBE = [round(x * 0.5, 1) for x in range(2, 121)]

def band_values(bl, p):
    """같은 (형상·가공부·코팅·날수) 에서 직경만 바꿔가며 나오는 값 집합."""
    vals = set()
    for d in PROBE:
        q = dict(p); q["dia"] = d
        try:
            v = bl.expect(**q)
            if v is not None: vals.add(round(v))
        except Exception: pass
    return vals

CELL = re.compile(rf'<c r="{PC}(\d+)"([^>]*?)(/>|>(.*?)</c>)', re.S)

def rewrite(xml, plan):
    hit = [0]
    def sub(m):
        r = int(m.group(1))
        if r not in plan or m.group(3) == "/>": return m.group(0)
        inner = m.group(4)
        if "<f" in (inner or ""): return m.group(0)
        new = re.sub(r'<v>[^<]*</v>', f'<v>{plan[r]}</v>', inner, count=1)
        if new == inner: return m.group(0)
        hit[0] += 1
        return f'<c r="{PC}{r}"{m.group(2)}>{new}</c>'
    return CELL.sub(sub, xml), hit[0]

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dir", default="raw/단가표/260922 단가 검수")
    ap.add_argument("--pricelist", default="raw/단가표/26.09 재연마정가표_정리.xlsx")
    ap.add_argument("--sheet", default="정가표(26.09)")
    ap.add_argument("--apply", action="store_true")
    a = ap.parse_args()

    pl = PriceList(a.pricelist, a.sheet); bl = Blocks(pl)
    if bl.missing: sys.exit(f"🔴 못 찾은 블록 {len(bl.missing)}개")
    print(f"모드 {'APPLY' if a.apply else 'DRY-RUN'}\n")

    tot = 0
    for fn in sorted(f for f in os.listdir(a.dir) if f.endswith("_완료.xlsx") and not f.startswith("~$")):
        path = os.path.join(a.dir, fn)
        ws = openpyxl.load_workbook(path, data_only=True)[SHEET_ITEMS]
        plan, detail = {}, []
        for r in range(ROW_START, ws.max_row + 1):
            code = ws.cell(r, COL_CODE).value
            nm   = ws.cell(r, COL_NAME).value
            pr   = ws.cell(r, COL_PRICE).value
            if not nm or not isinstance(pr, (int, float)): continue
            if code in IGNORE_CODES: continue
            p = parse_name(nm)
            if not p or "라핑" in p["shape"]: continue
            try: exp = bl.expect(**p)
            except Exception: exp = None
            if exp is None or round(pr) == round(exp): continue
            why = None
            if code in EXTRA_CODES:                       why = "지정 오류"
            elif p["blade"] in SKIP_BLADES:               why = None
            elif round(pr) in band_values(bl, p):         why = "구간 밀림"
            if not why: continue
            plan[r] = round(exp)
            detail.append((r, code, str(nm), round(pr), round(exp), why))
        tot += len(plan)
        print(f"  {fn}  → {len(plan)}건")
        for d in detail[:4]:
            print(f"     r{d[0]:<6} {str(d[1]):<13} {d[2][:32]:<34} {d[3]:>7,} → {d[4]:>7,}  [{d[5]}]")
        if len(detail) > 4: print(f"     … 외 {len(detail)-4}건")
        if not a.apply or not plan: continue

        zin = zipfile.ZipFile(path)
        xml = zin.read("xl/worksheets/sheet1.xml").decode("utf-8")
        dvb = xml.count("x14:dataValidation")
        new_xml, hits = rewrite(xml, plan)
        dva = new_xml.count("x14:dataValidation")
        if hits != len(plan) or dvb != dva:
            sys.exit(f"🔴 검증 실패: 치환 {hits}/{len(plan)} · 드롭다운 {dvb}→{dva}")
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
            for zi in zin.infolist():
                data = new_xml.encode("utf-8") if zi.filename == "xl/worksheets/sheet1.xml" \
                       else zin.read(zi.filename)
                zo = zipfile.ZipInfo(zi.filename, date_time=zi.date_time)
                zo.compress_type = zi.compress_type; zo.external_attr = zi.external_attr
                zout.writestr(zo, data)
        zin.close()
        with open(path, "wb") as f: f.write(buf.getvalue())   # 삭제 없이 덮어쓰기
        print(f"     ✅ 치환 {hits}/{len(plan)} · 드롭다운 {dvb}→{dva}")
    print(f"\n합계 {tot}건" + ("" if a.apply else "   (드라이런 — 실제 반영은 --apply)"))

if __name__ == "__main__":
    main()
