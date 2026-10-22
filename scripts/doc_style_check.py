#!/usr/bin/env python3
"""문서 문체 점검기 — 발표자료(pptx) · 보고서(docx) · 위키(md)에서 「AI 티」 패턴을 센다.

기준: wiki/standards/문서작성-AI티-점검표.md (2026-10-22)
쓰는 법:
    python scripts/doc_style_check.py 파일.pptx [파일.docx ...]
    python scripts/doc_style_check.py --strict 파일.pptx   # 기준 넘으면 종료 코드 1

고치지 않는다. 세기만 한다. 판단은 사람이 한다.
"""
import re, sys, os, unicodedata
from collections import Counter

# ---------- 1. 글 꺼내기 ----------
def load(path):
    """(위치, 종류, 글) 목록. 종류 = title(제목) / text(본문) / note(발표자 노트)"""
    ext = os.path.splitext(path)[1].lower()
    out = []
    if ext == ".pptx":
        from pptx import Presentation
        from pptx.enum.shapes import PP_PLACEHOLDER
        prs = Presentation(path)
        for i, sl in enumerate(prs.slides, 1):
            for sh in sl.shapes:
                kind = "text"
                if sh.is_placeholder and sh.placeholder_format.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
                    kind = "title"
                if sh.has_text_frame:
                    for p in sh.text_frame.paragraphs:
                        t = p.text.strip()
                        if t: out.append((f"{i}장", kind, t))
                if getattr(sh, "has_table", False) and sh.has_table:
                    for r in sh.table.rows:
                        for c in r.cells:
                            t = c.text.strip()
                            if t: out.append((f"{i}장 표", "text", t))
            if sl.has_notes_slide:
                t = sl.notes_slide.notes_text_frame.text.strip()
                if t: out.append((f"{i}장 노트", "note", t))
    elif ext == ".docx":
        import docx
        d = docx.Document(path)
        for j, p in enumerate(d.paragraphs, 1):
            t = p.text.strip()
            if not t: continue
            kind = "title" if p.style.name.lower().startswith(("heading", "title", "제목")) else "text"
            out.append((f"문단 {j}", kind, t))
        for ti, tb in enumerate(d.tables, 1):
            for r in tb.rows:
                for c in r.cells:
                    t = c.text.strip()
                    if t: out.append((f"표 {ti}", "text", t))
    else:
        for j, line in enumerate(open(path, encoding="utf-8"), 1):
            t = line.strip()
            if not t or t.startswith(("---", "|---")): continue
            kind = "title" if t.startswith("#") else "text"
            out.append((f"{j}행", kind, t.lstrip("# ")))
    return out

# ---------- 2. 패턴 ----------
JOSA = r"(은|는|이|가|을|를|의|에|에서|으로|로|와|과|도|만)"
P = {
    "대시(—, –)": (r"[—–]", "쉼표, 괄호, 문장 나누기로"),
    "가운데점 나열( · )": (r"\s·\s", "쉼표나 「및」으로"),
    "조사 앞 띄어쓰기": (r"(?<=[A-Za-z0-9\)])\s" + JOSA + r"(?=[\s,.」)]|$)", "붙여 쓰기 — 한글 맞춤법 제41항"),
    "「A가 아니라 B」 대구": (r"(가|이|은|는) 아니라", "사실 한 문장으로"),
    "이모지 · 판정 기호": (r"[✅❌⏸⚠🔴🟡🟢🔵⭐💡👉]", "「합격 / 보류」 같은 글자로"),
    "화살표(→)": (r"→", "「~에서 ~로」, 「~ 후」로"),
}
TRANSLATIONESE = ["을 통해", "를 통해", "에 있어서", "에 의해", "에 대한", "에 대해", "가지고 있", "되어지", "것이다", "하는 것이", "라고 할 수 있", "할 수 있을 것으로", "으로 보인다", "로 보인다"]
AI_PHRASE = ["결론적으로", "주목할 만", "시사하는 바", "중요한 역할", "핵심적", "다양한", "효과적으로", "획기적", "혁신적", "한눈에", "살펴보", "알아보", "에 그치지 않고", "뿐만 아니라", "궁극적으로", "본질적으로"]
SUMMARY_WORDS = ["축적", "확대", "해소", "우선", "고도화", "극대화", "강화", "혁신", "최적화", "효율화", "체계화", "구축 완료", "생성 완료"]
CLAUSE_COMMA = re.compile(r"(고|며|면서|서|면|지만|는데|은데|으나|하나|되나|이나|여|해|니|으로|로|도|다|음|함|됨|임)\s*,\s")  # 연결어미 · 종결 뒤 쉼표 = 절 잇기
TITLE_END_BAD = re.compile(r"(했음|었음|였음|있음|없음|됨|함|임|음)$")

LIMIT = {  # 기준값 — 점검표 §4
    "쉼표 문장 비율 %": 30,      # 사람 글 26 %, LLM 61 % (KatFishNet 2025) — 발표자료(pptx)에만 적용
    "쉼표 절 잇기 문장 %": 10,   # 「~하고, ~」 「~으로, ~」처럼 절과 절을 쉼표로 잇는 문장 — 모든 문서
    "대시(—, –)": 0,
    "가운데점 나열( · )": 3,
    "조사 앞 띄어쓰기": 0,
    "「A가 아니라 B」 대구": 0,
    "이모지 · 판정 기호": 0,
    "화살표(→)": 3,
    "번역투": 3,
    "AI 관용구": 2,
    "제목 요약어": 0,
    "제목 「~음 / ~함」 종결": 0,
    "제목 쉼표 두 마디": 2,
    "괄호 줄 비율 %": 25,
}

def check(items, include_notes=False):
    found = {k: [] for k in LIMIT}
    lines = [(loc, kind, t) for loc, kind, t in items if include_notes or kind != "note"]
    sent = comma = clause = paren = 0
    for loc, kind, t in lines:
        for name, (pat, _) in P.items():
            for m in re.finditer(pat, t):
                found[name].append((loc, t))
        for w in TRANSLATIONESE:
            if w in t: found["번역투"].append((loc, f"「{w}」 {t}"))
        for w in AI_PHRASE:
            if w in t: found["AI 관용구"].append((loc, f"「{w}」 {t}"))
        if "(" in t or "（" in t: paren += 1
        # 문장 단위 쉼표
        for s in re.split(r"(?<=[.!?。])\s+|\n", t):
            s = s.strip()
            if len(s) < 8: continue
            sent += 1
            if re.search(r"(?<!\d),|,(?!\d)|，", s): comma += 1  # 숫자 자릿점(10,744)은 제외
            if CLAUSE_COMMA.search(s + " "): clause += 1; found["쉼표 절 잇기 문장 %"].append((loc, s))
        if kind == "title":
            for w in SUMMARY_WORDS:
                if w in t: found["제목 요약어"].append((loc, f"「{w}」 {t}"))
            core = re.sub(r"[\s」』)\]]+$", "", t)
            if TITLE_END_BAD.search(core): found["제목 「~음 / ~함」 종결"].append((loc, t))
            if t.count(",") >= 1 and len(t) > 12: found["제목 쉼표 두 마디"].append((loc, t))
    stats = {
        "쉼표 문장 비율 %": round(100 * comma / sent, 1) if sent else 0,
        "쉼표 절 잇기 문장 %": round(100 * clause / sent, 1) if sent else 0,
        "괄호 줄 비율 %": round(100 * paren / len(lines), 1) if lines else 0,
    }
    return found, stats, len(lines), sent

def report(path, strict=False, include_notes=False):
    items = load(path)
    found, stats, n, sent = check(items, include_notes)
    print(f"\n=== {os.path.basename(path)} — 화면 글 {n}줄, 문장 {sent}개" + (" (노트 포함)" if include_notes else " (발표자 노트 제외)"))
    over = 0
    is_deck = path.lower().endswith(".pptx")
    for k, lim in LIMIT.items():
        v = stats[k] if k in stats else len(found[k])
        if k == "쉼표 문장 비율 %" and not is_deck:
            print(f"  [참고] {k}: {v} — 보고서 · 위키는 항목 나열 쉼표가 많아 판정하지 않음(절 잇기만 판정)")
            continue
        mark = "넘음" if v > lim else "통과"
        if v > lim: over += 1
        print(f"  [{mark}] {k}: {v} (기준 {lim})")
        if k in found and found[k] and v > lim:
            for loc, t in found[k][:5]:
                print(f"        - {loc}: {t[:90]}")
            if len(found[k]) > 5: print(f"        … 외 {len(found[k]) - 5}건")
    print(f"  → 기준 넘은 항목 {over}개")
    return over

if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    strict = "--strict" in sys.argv
    notes = "--notes" in sys.argv
    if not args:
        print(__doc__); sys.exit(0)
    total = sum(report(a, strict, notes) for a in args)
    sys.exit(1 if strict and total else 0)
