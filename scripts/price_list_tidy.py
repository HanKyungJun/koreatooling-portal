#!/usr/bin/env python3
# -*- coding: utf-8 -*-
r"""재연마 정가표 정리본 — 블록 매트릭스 → 한 줄 한 규격 표 (2026-10-27 Cowork)

한경준님(2026-10-27): 「정가표 자체를 정리」 · 형태 = 한 줄 한 규격 · 기준 = 지금 폴더의 판.
원본 정가표는 읽기만 한다(수정 금지 — 2026-09-02). 결과는 새 xlsx(저장소 밖 outputs/)로만 쓴다.

판독은 price_lookup.Lookup(= price_audit.Blocks + 일반코팅 별칭 보정 + 드릴 블록)과 같은 규칙.
직경을 0.05 mm 간격으로 짚어 「같은 셀을 읽는 구간」을 한 줄로 묶고, 그 셀 주소 · 축 라벨 · 블록 제목을 함께 적는다.

시트
  정가        한 줄 = 재질 · 형상 계열 · 가공부 · 코팅 · 날수 · 직경 구간 → 정가 (+ 셀 주소 · 블록 제목 · 비고)
  문제점      못 찾은 블록 · #VALUE! · 구간 라벨(이하 아님 · 빈틈) · 직경↑ 값↓ · 코팅 역전 · 제목 오기
  빈 조합     정가표에 값이 없는 조합(형상 계열 × 가공부 × 코팅 × 날수)
  읽는 법     기준 파일 · 규칙

실행: python scripts\price_list_tidy.py            → outputs\price_lookup\정가표_정리본_<시트>_<YYYYMMDD>.xlsx
"""
import collections, os, re, sys, warnings
warnings.filterwarnings('ignore')
HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import price_audit as pa
import price_lookup as L

FAM = [('평', '평'), ('볼 · 코너', '볼'), ('라핑평', '라핑평'), ('라핑볼 · 라핑코너', '라핑볼'), ('리머', '리머')]
PARTS = ['밑날', '밑옆날', '골수리', '옆날']
COATS = ['비코팅', '일반코팅', '고경도코팅']
BLADES = [(2, '2날'), (4, '3날 이상')]
STEP, DMAX = 0.05, 60.0
EXCEPT_E2 = pa.EXCEPTIONS_E2


def addr(r, c):
    return f'{get_column_letter(c)}{r}'


def main():
    lk = L.Lookup()
    pl, bl = lk.pl, lk.blocks
    flags = {}   # (r, c) 범위 → 비고(추정)
    # 라핑 평 밑골수리 고경도 — 제목이 「(일반코팅)」 으로 두 번 나오는 판: 오른쪽(두 번째) 블록을 고경도로 읽는다 [추정]
    k_hard = ('초경', 'L평', '밑골수리', '고경도')
    if k_hard not in bl.v:
        pos = pl.find(r'초경 라핑 평.*밑골수리', 1)
        if pos:
            bl.v[k_hard] = pl.vtable(*pos); bl.at[k_hard] = pos
            flags[pos] = '제목이 「(일반코팅)」 으로 왼쪽 블록과 같음 — 위치(오른쪽)와 값 크기로 고경도 블록으로 읽음 [추정 · 확인 필요]'
    # 블록 제목 위치 목록(셀 → 블록 제목 찾기용)
    tpos = sorted({p for p in bl.at.values()})
    for _pat in (r'^드릴/면취/NC 스포팅$', r'^초경_드릴/면취/NC 스포팅$', r'^HSS_드릴/면취/NC 스포팅$'):   # 26.10~ 재질별 제목(2026-10-29)
        if pl.find(_pat):
            tpos.append(pl.find(_pat))

    def title_of(r, c):
        cand = [(tr, tc) for tr, tc in tpos if tr <= r and tc <= c and r - tr <= 20 and c - tc <= 20]
        if not cand:
            return ''
        tr, tc = max(cand, key=lambda x: (x[1], x[0]))   # 가장 가까운 왼쪽 블록 → 그중 가장 아래
        return re.sub(r'\s+', ' ', str(pl.ws.cell(tr, tc).value or '')).strip(), (tr, tc)

    # 셀 · 축 라벨 기록용 계측
    last = {}
    og_gv, og_gh = pl.gv, pl.gh

    def gv(seg, col, dia):
        last.clear()
        for v, k, r in seg:
            if pa.in_band(v, k, dia):
                last.update(cell=(r, col), lab=(v, k), axis=[(a, b) for a, b, _ in seg]); break
        return og_gv(seg, col, dia)

    def gh(seg, row, dia):
        last.clear()
        for v, k, c in seg:
            if pa.in_band(v, k, dia):
                last.update(cell=(row, c), lab=(v, k), axis=[(a, b) for a, b, _ in seg]); break
        return og_gh(seg, row, dia)
    pl.gv, pl.gh = gv, gh

    rows, empty = [], []
    for fam_lab, shape in FAM:
        for part in PARTS:
            for coat in COATS:
                for blade, blab in BLADES:
                    n = int(round(DMAX / STEP))
                    for mat in ('초경', 'HSS'):
                        runs, prev = [], None
                        for i in range(1, n + 1):
                            d = round(i * STEP, 3)
                            try:
                                v = bl.expect(mat=mat, shape=shape, part=part, coat=coat, blade=blade, dia=d)
                            except Exception:
                                v = None
                            cell = last.get('cell') if v is not None else None
                            key = (v, cell)
                            if runs and key == prev:
                                runs[-1]['to'] = d
                            else:
                                runs.append(dict(fr=d, to=d, v=v, cell=cell, lab=last.get('lab') if v is not None else None))
                                prev = key
                        got = [r for r in runs if r['v'] is not None]
                        if not got:
                            empty.append((mat, fam_lab, part, coat, blab)); continue
                        lo = 0.0
                        for r in runs:
                            if r['v'] is None:
                                lo = r['to']; continue
                            lab = f"{r['lab'][0]:g}{r['lab'][1]}" if r['lab'] else ''
                            t = title_of(*r['cell'])
                            note = []
                            if blade == 4: note.append('3날 이상 = 4날 행')
                            fl = [m for p, m in flags.items() if t and t[1] == p]
                            note += fl
                            rows.append(dict(mat=mat, fam=fam_lab, part=part, coat=L.COAT_LABEL[coat], blade=blab,
                                             lo=lo, hi=r['to'], lab=lab, v=r['v'], cell=addr(*r['cell']), title=t[0] if t else '',
                                             note=' · '.join(note), est=bool(fl)))
                            lo = r['to']
    pl.gv, pl.gh = og_gv, og_gh
    # 드릴 ① 26.10~ — 「초경_ · HSS_드릴/면취/NC 스포팅」 재질별 블록(날경 「N이하」 · 비코팅 / 코팅 / 고경도) (2026-10-29)
    for _mat in ('초경', 'HSS'):
        mpos = pl.find(rf'^{_mat}_드릴/면취/NC 스포팅$')
        if not mpos:
            continue
        tr, tc = mpos
        hdr = {str(pl.ws.cell(tr + 1, c).value or '').strip(): c for c in range(tc, tc + 5)}
        lo, r = 0.0, tr + 2
        while r <= pl.ws.max_row:
            m = re.match(r'^\s*([\d.]+)\s*이하\s*$', str(pl.ws.cell(r, tc).value or ''))
            if not m:
                break
            hi = float(m.group(1))
            for lab_c, coat in (('비코팅', '비코팅'), ('코팅', '일반 코팅'), ('고경도', '고경도 코팅')):
                ci = hdr.get(lab_c)
                if not ci:
                    continue
                v = pl.ws.cell(r, ci).value
                rows.append(dict(mat=_mat, fam='드릴(면취 · NC 스포팅)', part='밑날', coat=coat, blade='2날',
                                 lo=lo, hi=hi, lab=str(pl.ws.cell(r, tc).value).strip(),
                                 v=v if isinstance(v, (int, float)) and v > 0 else None, cell=addr(r, ci),
                                 title=f'{_mat}_드릴/면취/NC 스포팅', note='날경 「N이하」 구간(앞 구간 초과 ~ N 이하)'))
            lo, r = hi, r + 1
    # 드릴 ② ~26.09 — 「드릴/면취/NC 스포팅」 블록(날경 a - b · 비코팅 / 코팅 · HSS/초경 동일 · 날수 구분 없음)
    dpos = None if pl.find(r'^초경_드릴/면취/NC 스포팅$') else pl.find(r'^드릴/면취/NC 스포팅$')
    if dpos:
        tr, tc = dpos; r = tr + 2
        while r <= pl.ws.max_row:
            got = False
            for c0 in (tc, tc + 3):
                m = re.match(r'^\s*([\d.]+)\s*-\s*([\d.]+)\s*$', str(pl.ws.cell(r, c0).value or ''))
                if m:
                    got = True
                    for ci, coat in ((c0 + 1, '비코팅'), (c0 + 2, '일반 코팅')):
                        v = pl.ws.cell(r, ci).value
                        rows.append(dict(mat='초경 · HSS 동일', fam='드릴(면취 · NC 스포팅)', part='밑날', coat=coat, blade='2날',
                                         lo=float(m.group(1)), hi=float(m.group(2)), lab=str(pl.ws.cell(r, c0).value).strip(),
                                         v=v if isinstance(v, (int, float)) else None, cell=addr(r, ci),
                                         title='드릴/면취/NC 스포팅',
                                         note='날경 a - b 구간(「이상 ~ 이하」) · 고경도 코팅 열 없음 — ERP 고경도 드릴은 코팅 × 약 1.07(10-27 확인)' if coat != '비코팅' else '날경 a - b 구간'))
            if not got:
                break
            r += 1

    # ── 문제점 ───────────────────────────────────────────
    probs = []
    for m in bl.missing:
        if '밑골수리' in m and '고경도' in m and k_hard in bl.v:
            continue   # 제목 오기로 따로 적음
        probs.append(('블록 없음', m, '', '이 제목의 블록을 정가표에서 찾지 못함 — 해당 조합은 정가표 값이 없음(정리본 · 조회기 모두 빈칸)'))
    for (r, c, v) in pl.errors:
        t = title_of(r, c)
        probs.append(('셀 오류', addr(r, c), t[0] if t else '', f'값이 {v}'))
    for p, msg in flags.items():
        probs.append(('제목 오기', addr(*p), str(pl.ws.cell(*p).value).strip(), msg))
    # 축 라벨: 이하가 아닌 것 · 숫자만 있는 칸
    seen = set()
    for key, (tr, tc) in bl.at.items():
        blk = bl.v.get(key) or bl.h.get(key)
        if not blk or not blk[0]:
            continue
        seg = blk[0]
        horiz = key in bl.h
        for j, (v, k, x) in enumerate(seg):
            is_last = (j == len(seg) - 1)
            r, c = (None, x) if horiz else (x, None)
            # 원래 칸 글자
            if horiz:
                # 가로형은 헤더 행을 다시 찾아 칸 글자를 본다
                for rr in list(range(tr + 1, tr + 4)) + list(range(tr - 1, max(0, tr - 9), -1)):
                    raw = pl.ws.cell(rr, x).value if rr >= 1 else None
                    b = pa.band(raw)
                    if b and b[0] == v:
                        r = rr; break
                c = x
            else:
                c = tc
            if r is None:
                continue
            raw = pl.ws.cell(r, c).value
            if (r, c) in seen:
                continue
            seen.add((r, c))
            s = str(raw).strip()
            bare = re.fullmatch(r'[\d.]+', s) is not None
            if s.endswith('미만'):
                why = '「미만」 — 경계값이 다음 구간과 어긋날 수 있음(정가표 구간은 「이하」 기준)'
            elif s.endswith('이상') and not is_last:
                why = '「이상」 이 마지막 칸이 아님 — 구간이 겹치거나 빔'
            elif bare and is_last:
                why = '마지막 칸이 숫자만 — 「이하」 로 읽혀 그보다 큰 직경이 빠짐(마지막은 「N이상」 이어야 함)'
            elif not bare and not s.endswith(('이하', '이상')):
                why = '라벨 형식 이상'
            else:
                why = None
            if why:
                probs.append(('구간 라벨', addr(r, c), title_of(r, c)[0] if title_of(r, c) else '', f'「{s}」 {why}'))
    # 직경↑ 값↓ · 코팅 역전 · 빈틈
    grp = collections.defaultdict(list)
    for x in rows:
        grp[(x['mat'], x['fam'], x['part'], x['coat'], x['blade'])].append(x)
    for k, xs in grp.items():
        xs.sort(key=lambda x: x['lo'])
        for a, b in zip(xs, xs[1:]):
            if a['v'] and b['v'] and b['v'] < a['v']:
                probs.append(('직경↑ 값↓', b['cell'], b['title'], f"{' · '.join(k)}: {a['lab']} {a['v']:,} → {b['lab']} {b['v']:,}"))
            if b['lo'] - a['hi'] > STEP + 1e-6 and k[1] != '드릴(면취 · NC 스포팅)':
                probs.append(('구간 빈틈', b['cell'], b['title'], f"{' · '.join(k)}: Ø{a['hi']:g} ~ Ø{b['lo']:g} 사이 값 없음"))
    idx = {}
    for x in rows:
        for d in (x['hi'],):
            idx[(x['mat'], x['fam'], x['part'], x['coat'], x['blade'], d)] = x
    for (mat, fam, part, coat, blade, d), x in idx.items():
        if coat != '비코팅':
            continue
        y = idx.get((mat, fam, part, '일반 코팅', blade, d)); z = idx.get((mat, fam, part, '고경도 코팅', blade, d))
        if y and x['v'] and y['v'] and x['v'] > y['v']:
            probs.append(('코팅 역전', y['cell'], y['title'], f"{mat} · {fam} · {part} · {blade} Ø{d:g}: 비코팅 {x['v']:,} > 일반 {y['v']:,}"))
        if y and z and y['v'] and z['v'] and y['v'] > z['v'] and (fam.split(' ')[0], mat, part, d) not in EXCEPT_E2:
            probs.append(('코팅 역전', z['cell'], z['title'], f"{mat} · {fam} · {part} · {blade} Ø{d:g}: 일반 {y['v']:,} > 고경도 {z['v']:,}"))

    # ── 쓰기 ─────────────────────────────────────────────
    out_dir = os.path.join(ROOT, 'outputs', 'price_lookup'); os.makedirs(out_dir, exist_ok=True)
    out = os.path.join(out_dir, f"정가표_정리본_{re.sub(r'[^0-9.]', '', pl.name)}_{pa.now_kst():%Y%m%d}.xlsx")
    wb = openpyxl.Workbook()
    FN = '맑은 고딕'
    ft, fs, fb, fbb = Font(name=FN, size=22, bold=True), Font(name=FN, size=15, bold=True), Font(name=FN, size=13), Font(name=FN, size=13, bold=True)
    H = PatternFill('solid', fgColor='D9E1F2'); Y = PatternFill('solid', fgColor='FFF2CC')

    def sheet(ws, title, sub, hdr, data, widths, money=()):
        ws['A1'] = title; ws['A1'].font = ft
        ws['A2'] = sub; ws['A2'].font = fb
        for c, h in enumerate(hdr, 1):
            x = ws.cell(4, c, h); x.font = fbb; x.fill = H; x.alignment = Alignment(wrap_text=True, vertical='center')
        for i, row in enumerate(data, 5):
            for c, v in enumerate(row, 1):
                x = ws.cell(i, c, v); x.font = fb
                if c in money and isinstance(v, (int, float)): x.number_format = '#,##0'
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = 'A5'
        return ws
    order = {f: i for i, (f, _) in enumerate(FAM)}; order['드릴(면취 · NC 스포팅)'] = 9
    data = []
    MO = {'초경': 0, 'HSS': 1}; CO = {'비코팅': 0, '일반 코팅': 1, '고경도 코팅': 2}; BO = {'2날': 0, '3날 이상': 1}
    for x in sorted(rows, key=lambda x: (order.get(x['fam'], 9), MO.get(x['mat'], 2), PARTS.index(x['part']) if x['part'] in PARTS else 0,
                                         CO.get(x['coat'], 9), BO.get(x['blade'], 9), x['lo'])):
        rng = f"Ø{x['lo']:g} 초과 ~ Ø{x['hi']:g} 이하" if x['fam'] != '드릴(면취 · NC 스포팅)' else f"Ø{x['lo']:g} ~ Ø{x['hi']:g}"
        if x['fam'] != '드릴(면취 · NC 스포팅)' and x['lo'] == 0:
            rng = f"Ø{x['hi']:g} 이하"
        if x['fam'] != '드릴(면취 · NC 스포팅)' and x['hi'] >= DMAX - 1e-9:
            rng = f"Ø{x['lo']:g} 초과"
        data.append([x['mat'], x['fam'], x['part'], x['coat'], x['blade'], rng, x['lab'], x['v'], x['cell'], x['title'], x['note']])
    ws = sheet(wb.active, '재연마 정가 — 한 줄 한 규격',
               f"원본: raw/단가표/{os.path.basename(lk.pl_path)} 「{pl.name}」(수정 {pa.datetime.datetime.fromtimestamp(os.path.getmtime(lk.pl_path), pa.KST):%Y-%m-%d %H:%M} KST) · 원본은 수정하지 않음 · 노란 줄 = 추정 판독",
               ['재질', '형상 계열', '가공부', '코팅', '날수', '직경 구간', '정가표 축 라벨', '정가(원)', '원본 셀', '원본 블록 제목', '비고'],
               data, [10, 18, 9, 12, 10, 24, 12, 12, 9, 46, 40], money=(8,))
    ws.title = '정가'
    for i, x in enumerate(data, 5):
        if '추정' in (x[10] or ''):
            for c in range(1, 12): ws.cell(i, c).fill = Y
    sheet(wb.create_sheet('문제점'), '정가표 문제점 — 원본에서 고칠 곳', '원본 셀 주소 기준. Cowork 는 원본을 고치지 않는다(한경준님이 판단 · 정정)',
          ['구분', '원본 셀', '블록 제목', '내용'], sorted(probs, key=lambda p: (p[0], p[1])), [12, 10, 46, 80])
    sheet(wb.create_sheet('빈 조합'), '정가표에 값이 없는 조합', '형상 계열 × 가공부 × 코팅 × 날수 기준. 원래 없는 가공(예: 리머 옆날)도 섞여 있다',
          ['재질', '형상 계열', '가공부', '코팅', '날수'], [[a, b, c, L.COAT_LABEL[d], e] for a, b, c, d, e in empty], [10, 18, 10, 12, 10])
    rd = wb.create_sheet('읽는 법')
    rd['A1'] = '읽는 법'; rd['A1'].font = ft
    lines = ['· 한 줄 = 한 규격 구간. 「Ø4 초과 ~ Ø6 이하」 처럼 위쪽 경계를 포함한다(정가표 구간은 「이하」 기준)',
             '· 「3날 이상」 = 정가표의 4날(3,4,6날) 행. 형상 계열 「볼 · 코너」 는 같은 블록을 쓴다',
             '· 가공부: 밑옆날 = 밑골수리 = 밑외경(같은 가공). 골수리 · 옆날은 외경연삭 표를 쓴다(정가표 블록 주석 규칙)',
             '· 라핑 「라핑」 · 「라핑평」 은 같은 값. 초경 라핑 밑날 「= 초경 밑날 × 1.1」 같은 주석 규칙은 원본 블록 값 그대로 옮긴 것',
             '· 원본 셀 · 블록 제목으로 원본 정가표의 그 칸을 바로 찾을 수 있다',
             '· 이 파일은 원본에서 다시 만든다: python scripts\\price_list_tidy.py (원본이 바뀌면 다시 실행)',
             f'· 생성 {pa.now_kst():%Y-%m-%d %H:%M} KST · 판독 규칙: wiki/standards/재연마-정가표-읽는법.md']
    for i, s in enumerate(lines, 3):
        rd.cell(i, 1, s).font = fb
    rd.column_dimensions['A'].width = 120
    wb.save(out)
    print(f'정리본 저장: {out}')
    print(f'  정가 {len(data)}줄 · 문제점 {len(probs)}건 · 빈 조합 {len(empty)}개 · 못 찾은 블록 {len(bl.missing)}개')
    c = collections.Counter(p[0] for p in probs); print('  문제점 구분:', dict(c))


if __name__ == '__main__':
    main()
