"""ERP 정가 재검수 — 내려받은(또는 올릴) 품목 파일을 정가표 · 사업장 사슬 · 기존 값과 대조 (2026-10-29 Cowork)

실행(PC, cnc-wiki 폴더에서):
    python scripts\\price_erp_check.py                      # raw/단가표 의 가장 새 폴더 ↔ 그 앞 「품목정보조회」 폴더
    python scripts\\price_erp_check.py 261031 261027         # 새 폴더 · 기준 폴더를 직접 지정
    python scripts\\price_erp_check.py 261031 261027 --drill  # 드릴(소분류 · 품목명 「드릴」)만

읽는 파일(폴더 안, 사업장 이름으로 시작): {코툴|코고|창성}_품목정보조회_*.xlsx 또는 {코툴|코고|창성}_품목일괄등록_*.xlsx
  - 품목정보조회: 1 품목코드 · 2 품목명 · 6 정가 · 7 매출할인율 · 8 매출단가 · 9 매입할인율 · 10 매입단가 (3행부터)
  - 품목일괄등록: 2 품목코드 · 3 품목명 · 28 정가 · 29 매출할인율 · 30 매출단가 · 31 매입할인율 · 32 매입단가 (3행부터)

검사:
  ① 정가 ↔ 정가표 계산(scripts/price_lookup.py 와 같은 판정 — 블록 있는 형상 · 2날 밑날 드릴)
  ② 단가 = 정가 × 할인율 — 10원 올림 · 버림 · 반올림 중 어느 것인지, 어느 것도 아니면 「계산 다름」. 빈칸 · 0 도 표시
  ③ 사업장 사슬 — 코툴 매입 = 코고 매출 · 코고 매입 = 창성 매입 (창성 가격은 창성 파일 「매입」 칸 — 2026-10-29 정정)
  ④ 기존 대비 — 코툴 · 코고 = 매출단가, 창성 = 매입단가. 내려간 품목 목록(HSS 드릴은 인하가 의도 — decisions 2026-10-29 (6))
결과: wiki/_private/ERP정가_재검수_{새폴더}_vs_{기준폴더}.xlsx (실거래 단가 — 대외비, git 제외 폴더)
정가표 · ERP 파일은 읽기만 한다.
"""
import collections
import glob
import math
import os
import re
import sys
import warnings

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L

warnings.filterwarnings('ignore')
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
PRICE_DIR = os.path.join(ROOT, 'raw', '단가표')
OUT_DIR = os.path.join(ROOT, 'wiki', '_private')
SITES = ('코툴', '코고', '창성')
FMT = {  # 1-based 열: 코드, 이름, 정가, 매출율, 매출단가, 매입율, 매입단가, 첫 데이터 행
    '품목정보조회': (1, 2, 6, 7, 8, 9, 10, 3),
    '품목일괄등록': (2, 3, 28, 29, 30, 31, 32, 3),
}


def num(x):
    try:
        v = float(x)
        return v if not math.isnan(v) else None
    except (TypeError, ValueError):
        return None


def ru(x): return int(math.ceil(round(x, 6) / 10) * 10)
def rd(x): return int(math.floor(round(x, 6) / 10) * 10)
def rn(x): return int(math.floor(round(x, 6) / 10 + 0.5) * 10)


def base_code(code):
    """사업장 접두(코툴 K · 창성 P) 를 떼어 코고 코드로 맞춘다."""
    c = str(code or '').strip()
    return c[1:] if c[:1] in ('K', 'k', 'P', 'p') and len(c) > 1 and c[1].isdigit() else c


def parse_name(n):
    """「2날/드릴/6mm/초경/밑날/비코팅」 → (blade, shape, dia, mat, part, coat) 또는 None."""
    p = [s.strip() for s in str(n or '').replace('//', '/').split('/')]
    if len(p) < 6:
        return None
    b = re.match(r'^(\d+)', p[0]); d = re.match(r'^([\d.]+)', p[2])
    if not b or not d:
        return None
    mat = '초경' if '초경' in p[3] else ('HSS' if 'hss' in p[3].lower() else p[3])
    coat = p[5].replace(' ', '')
    coat = '고경도코팅' if '고경도' in coat else ('비코팅' if '비코팅' in coat else ('일반코팅' if '코팅' in coat else coat))
    shape = '라핑평' if p[1] == '라핑' else p[1]
    return int(b.group(1)), shape, float(d.group(1)), mat, p[4].replace(' ', ''), coat


def find_files(folder):
    out = {}
    for s in SITES:
        for kind in FMT:
            fs = sorted(f for f in glob.glob(os.path.join(folder, f'{s}_{kind}_*.xlsx')) if not os.path.basename(f).startswith('~$'))
            if fs:
                out[s] = (fs[-1], kind)
                break
    return out


def load(path, kind, drill_only=False):
    cc, cn, cp, csr, csu, cbr, cbu, r0 = FMT[kind]
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb['품목일괄등록'] if kind == '품목일괄등록' and '품목일괄등록' in wb.sheetnames else wb.worksheets[0]
    rows = {}
    for i, r in enumerate(ws.iter_rows(values_only=True), start=1):
        if i < r0 or not r or len(r) < cn or not r[cc - 1]:
            continue
        name = str(r[cn - 1] or '').strip()
        if drill_only and '드릴' not in name:
            continue
        g = lambda c: r[c - 1] if len(r) >= c else None
        rows[str(r[cc - 1]).strip()] = dict(name=name, p=num(g(cp)), sr=num(g(csr)), su=num(g(csu)), br=num(g(cbr)), bu=num(g(cbu)))
    return rows


def unit_check(p, r, u):
    if p is None or p <= 0:
        return '정가 없음'
    if r is None:
        return '할인율 빈칸'
    if u is None or u == 0:
        return '단가 빈칸'
    x = p * r
    if u == ru(x): return '올림'
    if u == rd(x): return '버림'
    if u == rn(x): return '반올림'
    return '계산 다름'


def main():
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    drill_only = '--drill' in sys.argv
    dirs = sorted(d for d in glob.glob(os.path.join(PRICE_DIR, '*')) if os.path.isdir(d) and find_files(d))
    if not dirs:
        sys.exit('raw/단가표 아래에 품목정보조회 · 품목일괄등록 폴더가 없음')
    new_dir = os.path.join(PRICE_DIR, args[0]) if args else dirs[-1]
    if len(args) > 1:
        old_dir = os.path.join(PRICE_DIR, args[1])
    else:
        cand = [d for d in dirs if os.path.basename(d) < os.path.basename(new_dir) and glob.glob(os.path.join(d, '*_품목정보조회_*.xlsx'))]
        old_dir = cand[-1] if cand else None
    nf, of = find_files(new_dir), (find_files(old_dir) if old_dir else {})
    print(f'새 폴더: {os.path.basename(new_dir)} {[(s, os.path.basename(f)) for s, (f, _) in nf.items()]}')
    print(f'기준 폴더: {os.path.basename(old_dir) if old_dir else "(없음)"}')
    new = {s: load(f, k, drill_only) for s, (f, k) in nf.items()}
    old = {s: load(f, k, drill_only) for s, (f, k) in of.items()}

    # ① 정가표 계산 — price_lookup 엔진(같은 폴더). 못 불러오면 이 검사만 건너뛴다
    lk = None
    try:
        sys.path.insert(0, HERE)
        import price_lookup as pl
        lk = pl.Lookup(PRICE_DIR)
        pl_name = f'{os.path.basename(lk.pl_path)} 「{lk.pl.name}」'
    except Exception as e:  # noqa: BLE001
        pl_name = f'(정가표 엔진 불러오기 실패 — ① 건너뜀: {e})'
    print('정가표:', pl_name)

    def calc(q):
        if not lk or not q:
            return None
        b, s, d, m, p, c = q
        try:
            if s == '드릴':
                return lk.drill_price(d, c, m) if (b == 2 and p == '밑날') else None
            if s in ('평', '볼', '코너', '라핑평', '라핑볼', '라핑코너', '리머'):
                return lk.blocks.expect(mat=m, shape=s, part=p, coat=c, blade=b, dia=d)
        except Exception:  # noqa: BLE001
            return None
        return None

    res = collections.OrderedDict()
    summ = []
    # ① · ②
    s1, s2 = [], []
    for s, rows in new.items():
        c1, c2 = collections.Counter(), collections.Counter()
        for code, d in rows.items():
            e = calc(parse_name(d['name'])) if s != '창성' else None   # 창성 정가 칸은 매입 기준이라 정가표와 다를 수 있음(한경준님 10-29) — 대조 생략
            if s == '창성':
                c1['대조 생략(창성 = 매입 기준)'] += 1
            elif e is None:
                c1['정가표 대상 아님(블록 없음 · 표 밖)'] += 1
            elif d['p'] is not None and abs(d['p'] - e) < 0.5:
                c1['정가 일치'] += 1
            else:
                c1['정가 불일치'] += 1
                s1.append([s, code, d['name'], d['p'], e, (d['p'] or 0) - e])
            for lab, rk, uk in (('매출', 'sr', 'su'), ('매입', 'br', 'bu')):
                st = unit_check(d['p'], d[rk], d[uk])
                c2[f'{lab} {st}'] += 1
                if st in ('계산 다름', '할인율 빈칸', '단가 빈칸', '정가 없음'):
                    note = '코고 매입은 정가 · 매출만 바꾸고 안 고침 — 창성 매입과 같으면 정상(10-29)' if (s == '코고' and lab == '매입') else ''
                    s2.append([s, lab, code, d['name'], d['p'], d[rk], d[uk], st, note])
        summ.append((f'{s} 정가표 대조', dict(c1)))
        summ.append((f'{s} 단가 계산', dict(c2)))
    # ③ 사슬
    s3 = []
    kt, kg, cs = new.get('코툴', {}), new.get('코고', {}), new.get('창성', {})
    gidx = {base_code(c): d for c, d in kg.items()}
    cidx = {base_code(c): d for c, d in cs.items()}
    c3 = collections.Counter()
    for code, d in kt.items():
        g = gidx.get(base_code(code))
        if not g:
            c3['코툴 → 코고 짝 없음'] += 1; continue
        if d['bu'] == g['su']:
            c3['코툴 매입 = 코고 매출'] += 1
        else:
            c3['코툴 매입 ≠ 코고 매출'] += 1; s3.append(['코툴 매입 ↔ 코고 매출', code, d['name'], d['bu'], g['su']])
    for code, g in kg.items():
        c = cidx.get(base_code(code))
        if not c:
            c3['코고 → 창성 짝 없음'] += 1; continue
        if g['bu'] is None:
            c3['코고 매입 빈칸(창성 있음)'] += 1; s3.append(['코고 매입 빈칸', code, g['name'], None, c['bu']]); continue
        if g['bu'] == c['bu']:
            c3['코고 매입 = 창성 매입'] += 1
        else:
            c3['코고 매입 ≠ 창성 매입'] += 1; s3.append(['코고 매입 ↔ 창성 매입', code, g['name'], g['bu'], c['bu']])
    for s, rows in (('코툴', kt), ('코고', kg)):
        for code, d in rows.items():
            if d['su'] and d['bu'] and d['su'] < d['bu']:
                c3[f'{s} 역마진(매출 < 매입)'] += 1; s3.append([f'{s} 역마진', code, d['name'], d['su'], d['bu']])
    summ.append(('사업장 사슬', dict(c3)))
    # ④ 기존 대비
    s4, grp = [], collections.defaultdict(list)
    for s, field in (('코툴', 'su'), ('코고', 'su'), ('창성', 'bu')):
        for code, d in new.get(s, {}).items():
            o = old.get(s, {}).get(code)
            if not o or not d[field] or not o[field]:
                continue
            x = d[field] / o[field] - 1
            q = parse_name(d['name'])
            key = (s, q[3] if q else '?', q[1] if q else '?')
            grp[key].append(x)
            if x < 0:
                intent = 'HSS 드릴 인하(의도)' if q and q[1] == '드릴' and q[3] == 'HSS' else ''
                s4.append([s, '매출' if field == 'su' else '매입', code, d['name'], o['p'], o[field], d['p'], d[field], x, intent])
    c4 = {f'{k[0]} {k[1]} {k[2]}': f'{len(v)}건 평균 {sum(v) / len(v):+.1%} · 내려감 {sum(1 for y in v if y < 0)}' for k, v in sorted(grp.items())}
    summ.append(('기존 대비(형상별)', c4))

    # ── 저장 ──
    os.makedirs(OUT_DIR, exist_ok=True)
    out = os.path.join(OUT_DIR, f'ERP정가_재검수_{os.path.basename(new_dir)}_vs_{os.path.basename(old_dir) if old_dir else "없음"}' + ('_드릴' if drill_only else '') + '.xlsx')
    wb = openpyxl.Workbook()
    F = '맑은 고딕'; hf = PatternFill('solid', fgColor='D9E1F2')
    ws = wb.active; ws.title = '요약'
    ws['A1'] = 'ERP 정가 재검수'; ws['A1'].font = Font(name=F, size=22, bold=True)
    ws['A2'] = f'새 {os.path.basename(new_dir)} ↔ 기준 {os.path.basename(old_dir) if old_dir else "없음"} · 정가표 {pl_name} · 대외비(실거래 단가)'
    ws['A2'].font = Font(name=F, size=11, color='555555')
    r = 4
    for title, d in summ:
        ws.cell(r, 1, title).font = Font(name=F, size=15, bold=True); r += 1
        for k, v in (d.items() if d else [('(없음)', '')]):
            ws.cell(r, 1, k).font = Font(name=F, size=13); ws.cell(r, 2, v).font = Font(name=F, size=13); r += 1
        r += 1
    ws.column_dimensions['A'].width = 40; ws.column_dimensions['B'].width = 40

    def sheet(name, head, rows, money=(), pct=()):
        sh = wb.create_sheet(name)
        sh.append(head)
        for c in sh[1]:
            c.font = Font(name=F, bold=True, size=11); c.fill = hf; c.alignment = Alignment(horizontal='center', wrap_text=True)
        for row in rows:
            sh.append(row)
        for row in sh.iter_rows(min_row=2):
            for c in row:
                c.font = Font(name=F, size=10)
                if c.column in money: c.number_format = '#,##0'
                if c.column in pct: c.number_format = '0.0%;[Red]-0.0%'
        for i in range(1, len(head) + 1):
            sh.column_dimensions[L(i)].width = 34 if '품목명' in head[i - 1] else 13
        sh.freeze_panes = 'A2'
        if rows:
            sh.auto_filter.ref = f'A1:{L(len(head))}{len(rows) + 1}'
    sheet('① 정가 불일치', ['사업장', '품목코드', '품목명', 'ERP 정가 (원)', '정가표 계산 (원)', '차이 (원)'], s1, money=(4, 5, 6))
    sheet('② 단가 확인', ['사업장', '구분', '품목코드', '품목명', '정가 (원)', '할인율', '단가 (원)', '판정', '비고'], s2, money=(5, 7))
    sheet('③ 사슬', ['구분', '품목코드', '품목명', '값 1 (원)', '값 2 (원)'], s3, money=(4, 5))
    sheet('④ 내려간 품목', ['사업장', '구분', '품목코드', '품목명', '기존 정가 (원)', '기존 단가 (원)', '새 정가 (원)', '새 단가 (원)', '차이 (%)', '비고'],
          sorted(s4, key=lambda x: x[8]), money=(5, 6, 7, 8), pct=(9,))
    wb.save(out)
    print('저장:', out)
    for title, d in summ:
        print(f'  {title}: {d}')


if __name__ == '__main__':
    main()
