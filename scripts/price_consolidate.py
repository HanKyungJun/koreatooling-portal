#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
3사 단가 파일 → 통합 검수 파일 (정정 내역 정렬 포함)

왜: 창성·코고·코툴 3파일을 매입처만 다른 같은 구조로 따로 보면 같은 품목의
    어긋남을 못 본다. 정가표 대조 결과를 붙여 한 파일로 모으고, 정정 대상을
    「라핑(이번 개정 본체)」과 「라핑 아닌 것(확인 필요)」으로 갈라 정렬한다.

⚠️ 읽기만 한다. 원본 3파일도 정가표도 수정하지 않는다.

사용: python scripts/price_consolidate.py
"""
import argparse, collections, datetime, os, re, sys
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from price_audit import PriceList, Blocks, parse_name, SHEET_ITEMS, ROW_START, \
                        COL_CODE, COL_NAME, COL_PRICE

MG = "맑은 고딕"
F_T  = Font(name=MG, size=22, bold=True)
F_S  = Font(name=MG, size=15, bold=True)
F_B  = Font(name=MG, size=13)
F_BH = Font(name=MG, size=13, bold=True)
FH = PatternFill("solid", fgColor="DDEBF7")
FG = PatternFill("solid", fgColor="F2F2F2")
FW = PatternFill("solid", fgColor="FFF2CC")
FR = PatternFill("solid", fgColor="FCE4E4")
FB = PatternFill("solid", fgColor="E2EFDA")
_t = Side(style="thin", color="BFBFBF"); BD = Border(_t, _t, _t, _t)

SHAPE_ORDER = ["평","코너","볼","라핑평","라핑","라핑코너","라핑볼",
               "날붙이초경","날붙이초경볼","리머","드릴","챔퍼","절단","원통연삭","테이퍼"]
PART_ORDER  = ["밑날","옆날","골수리","외경연삭","밑옆날","밑골수리","밑외경"]
COAT_ORDER  = ["비코팅","일반코팅","고경도코팅","코팅"]
def _k(lst, v): return (lst.index(v) if v in lst else 99, str(v))
def sort_key(x):
    return (x["vend"], _k(SHAPE_ORDER, x["shape"]), _k(PART_ORDER, x["part"]),
            _k(COAT_ORDER, x["coat"]), x["blade"], x["dia"])

VEND = {"창성": 1, "코고": 2, "코툴": 3}

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dir", default="raw/단가표/260922 단가 검수")
    ap.add_argument("--pricelist", default="raw/단가표/26.09 재연마정가표_정리.xlsx")
    ap.add_argument("--sheet", default="정가표(26.09)")
    ap.add_argument("--out", default=None)
    a = ap.parse_args()

    pl = PriceList(a.pricelist, a.sheet)
    bl = Blocks(pl)
    if bl.missing:
        print(f"🔴 못 찾은 블록 {len(bl.missing)}개 — 먼저 제목 패턴을 맞추세요")
        for m in bl.missing: print("   ", m)
        sys.exit(1)

    files = sorted(f for f in os.listdir(a.dir) if f.endswith(".xlsx") and not f.startswith("~$"))
    recs = []
    for fn in files:
        vm = re.search(r"_([가-힣]+)\.xlsx$", fn)
        vend = vm.group(1) if vm else fn[:8]
        ws = openpyxl.load_workbook(os.path.join(a.dir, fn), data_only=True)[SHEET_ITEMS]
        for r in range(ROW_START, ws.max_row + 1):
            code = ws.cell(r, COL_CODE).value
            nm   = ws.cell(r, COL_NAME).value
            pr   = ws.cell(r, COL_PRICE).value
            if not code and not nm: continue
            p = parse_name(nm) if nm else None
            exp = None
            if p:
                try: exp = bl.expect(**p)
                except Exception: exp = None
            cur = pr if isinstance(pr, (int, float)) else None
            if p is None:                      verdict = "형식오류"
            elif cur is None:                  verdict = "정가 결측"
            elif exp is None:                  verdict = "미대조"
            elif round(cur) == round(exp):     verdict = "일치"
            else:                              verdict = "불일치"
            lap = bool(p and "라핑" in p["shape"])
            recs.append(dict(vend=vend, code=code or "", nm=str(nm or ""), row=r,
                             blade=p["blade"] if p else 0, shape=p["shape"] if p else "",
                             dia=p["dia"] if p else 0.0, mat=p["mat"] if p else "",
                             part=p["part"] if p else "", coat=p["coat"] if p else "",
                             cur=cur, exp=round(exp) if exp is not None else None,
                             verdict=verdict, lap=lap))
    print(f"총 {len(recs)}건 · 파일 {len(files)}개")
    cnt = collections.Counter(x["verdict"] for x in recs)
    for k, v in cnt.most_common(): print(f"   {k}: {v}")

    fix_lap  = sorted([x for x in recs if x["verdict"]=="불일치" and x["lap"]], key=sort_key)
    fix_else = sorted([x for x in recs if x["verdict"]=="불일치" and not x["lap"]], key=sort_key)
    uncov    = sorted([x for x in recs if x["verdict"]=="미대조"], key=sort_key)
    print(f"\n   → 라핑 정정대상 {len(fix_lap)} · 비라핑 불일치 {len(fix_else)} · 미대조 {len(uncov)}")

    now = datetime.datetime.utcnow() + datetime.timedelta(hours=9)
    out = a.out or f"outputs/단가검수_통합_{now:%Y%m%d}.xlsx"
    os.makedirs(os.path.dirname(out) or ".", exist_ok=True)
    wb = openpyxl.Workbook()

    HDR = ["매입처","품목코드","품목명","날수","형상","직경","재질","가공부","코팅",
           "현재 단가","정가표","차액","증감율(%)","원본 행"]
    def sheet(name, title, note, data, fill=None, with_exp=True):
        ws = wb.create_sheet(name)
        ws["A1"] = title; ws["A1"].font = F_T
        ws["A2"] = note;  ws["A2"].font = F_B
        for i, h in enumerate(HDR, 1):
            c = ws.cell(4, i, h); c.font = F_BH; c.fill = FH; c.border = BD
            c.alignment = Alignment(horizontal="center", wrap_text=True)
        for j, x in enumerate(data, 5):
            dif = (x["exp"] - x["cur"]) if (x["exp"] is not None and x["cur"] is not None) else None
            rate = (dif / x["cur"] * 100) if (dif is not None and x["cur"]) else None
            vals = [x["vend"], x["code"], x["nm"], x["blade"], x["shape"], x["dia"],
                    x["mat"], x["part"], x["coat"], x["cur"],
                    x["exp"] if with_exp else "", dif if with_exp else "",
                    round(rate,1) if (with_exp and rate is not None) else "", x["row"]]
            for i, v in enumerate(vals, 1):
                c = ws.cell(j, i, v); c.font = F_B; c.border = BD
                if i in (10, 11, 12): c.number_format = "#,##0"
                if fill and i in (10, 11): c.fill = fill
        for col, w in zip(HDR, (8,15,40,7,11,8,7,12,11,12,12,11,11,9)):
            ws.column_dimensions[get_column_letter(HDR.index(col)+1)].width = w
        ws.freeze_panes = "A5"
        if data: ws.auto_filter.ref = f"A4:N{len(data)+4}"
        return ws

    ws = wb.active; ws.title = "요약"
    ws["A1"] = "2026-09-22 단가 검수 — 3사 통합"; ws["A1"].font = F_T
    info = [
        ("생성", f"{now:%Y-%m-%d %H:%M} KST", "scripts/price_consolidate.py"),
        ("품목 파일", " · ".join(files), f"{len(recs)}건"),
        ("정가표", f"{os.path.basename(a.pricelist)} / {a.sheet}", "읽기 전용 — 수정하지 않음"),
        ("", "", ""),
        ("판정", "", ""),
        ("✅ 일치", cnt.get("일치", 0), "정가표와 같음 — 손댈 것 없음"),
        ("🟢 라핑 불일치", len(fix_lap), "이번 개정 본체. 정가표 값으로 정정 → 시트 「정정_라핑」"),
        ("⚠️ 비라핑 불일치", len(fix_else), "의도된 값일 수 있어 일괄 적용 안 함 → 시트 「확인_비라핑」"),
        ("⬜ 미대조", len(uncov), "정가표에 해당 블록이 없음 → 시트 「미대조」"),
        ("🔴 정가 결측", cnt.get("정가 결측", 0), "단가가 비어 있음"),
        ("🔴 형식오류", cnt.get("형식오류", 0), "품목명이 6필드가 아님 — 대조 자체가 안 됨"),
        ("", "", ""),
        ("정렬 기준", "매입처 → 형상 → 가공부 → 코팅 → 날수 → 직경", "형상: 평·코너·볼 → 라핑평·라핑·라핑코너·라핑볼 → 그 외"),
        ("", "", ""),
        ("⚠️ 이 파일은 스냅샷입니다", "정가표나 품목파일이 바뀌면 다시 돌리세요", "python scripts/price_consolidate.py"),
    ]
    r = 3
    for k, v, n in info:
        ws.cell(r,1,k); ws.cell(r,2,v); ws.cell(r,3,n)
        f = F_S if (k and v=="" and n=="") else F_B
        for c in (1,2,3):
            ws.cell(r,c).font = f
            ws.cell(r,c).alignment = Alignment(vertical="center", wrap_text=True)
        r += 1
    for col, w in zip("ABC", (22, 46, 60)): ws.column_dimensions[col].width = w

    sheet("정정_라핑", f"🟢 라핑 정정 대상 {len(fix_lap)}건",
          "정가표(26.09) 값으로 바꾸면 되는 것들입니다. 차액·증감율은 현재값 기준.", fix_lap, FB)
    sheet("확인_비라핑", f"⚠️ 라핑 아닌 불일치 {len(fix_else)}건 — 정정 전 확인",
          "라핑 개정과 무관한데 정가표와 다릅니다. 의도된 값일 수 있어 자동 정정하지 않았습니다.", fix_else, FR)
    sheet("미대조", f"⬜ 미대조 {len(uncov)}건",
          "정가표에 해당 블록이 없어 대조가 불가능합니다(드릴·챔퍼 등).", uncov, FW, with_exp=False)

    ws = wb.create_sheet("미대조_집계")
    ws["A1"] = "미대조 계열 집계"; ws["A1"].font = F_T
    for i, h in enumerate(["재질","형상","가공부","건수"], 1):
        c = ws.cell(3, i, h); c.font = F_BH; c.fill = FH; c.border = BD
    agg = collections.Counter((x["mat"], x["shape"], x["part"]) for x in uncov)
    for j, ((m, sh, pt), n) in enumerate(agg.most_common(), 4):
        for i, v in enumerate([m, sh, pt, n], 1):
            c = ws.cell(j, i, v); c.font = F_B; c.border = BD
    for col, w in zip("ABCD", (9, 14, 14, 9)): ws.column_dimensions[col].width = w
    ws.freeze_panes = "A4"

    wb.save(out)
    print(f"\n저장: {out}")

if __name__ == "__main__":
    main()
